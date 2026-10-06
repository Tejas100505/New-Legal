"""
Report Generation Agent for the E-Discovery platform.

Exports genuine, evidence-backed legal reports in PDF (via reportlab), DOCX (via python-docx),
Markdown, CSV, and JSON across four core legal report types:
- executive_summary
- review_summary
- privilege_log
- risk_assessment
"""

from __future__ import annotations

import base64
import csv
import html
import io
import json
from datetime import datetime, timezone
from typing import Any, Dict, List, Optional
from pydantic import BaseModel, Field
from loguru import logger

from app.agents.base import BaseAgent, AgentException

# ReportLab imports
try:
    from reportlab.lib.pagesizes import letter
    from reportlab.platypus import (
        SimpleDocTemplate,
        Paragraph,
        Spacer,
        Table,
        TableStyle,
        KeepTogether,
        PageBreak,
        HRFlowable,
    )
    from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
    from reportlab.lib import colors
    from reportlab.pdfgen import canvas
    REPORTLAB_AVAILABLE = True
except ImportError:
    REPORTLAB_AVAILABLE = False

# python-docx imports
try:
    import docx
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.table import WD_TABLE_ALIGNMENT
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False


class ReportInput(BaseModel):
    document_id: str
    report_type: str = "review_summary"  # executive_summary, review_summary, privilege_log, risk_assessment
    format: str = "markdown"  # markdown, json, csv, pdf, docx
    case_info: Dict[str, Any] = Field(default_factory=dict)
    corpus_stats: Dict[str, Any] = Field(default_factory=dict)
    executive_summary: Optional[str] = None
    executive_sections: Dict[str, Any] = Field(default_factory=dict)
    summary_metadata: Dict[str, Any] = Field(default_factory=dict)
    top_entities: List[Dict[str, Any]] = Field(default_factory=list)
    findings: List[Dict[str, Any]] = Field(default_factory=list)
    timeline: List[Dict[str, Any]] = Field(default_factory=list)
    privilege_log: List[Dict[str, Any]] = Field(default_factory=list)
    risks: List[Dict[str, Any]] = Field(default_factory=list)
    recommendations: List[str] = Field(default_factory=list)
    review_data: Dict[str, Any] = Field(default_factory=dict)


class ReportOutput(BaseModel):
    document_id: str
    report_type: str
    format: str
    content: str  # Raw text for MD/CSV/JSON, Base64-encoded for PDF/DOCX
    confidence_score: float = 1.0


if REPORTLAB_AVAILABLE:
    class NumberedCanvas(canvas.Canvas):
        """Two-pass canvas for Page X of Y and professional running headers/footers."""

        def __init__(self, *args, **kwargs):
            super().__init__(*args, **kwargs)
            self._saved_page_states = []

        def showPage(self):
            self._saved_page_states.append(dict(self.__dict__))
            self._startPage()

        def save(self):
            num_pages = len(self._saved_page_states)
            for state in self._saved_page_states:
                self.__dict__.update(state)
                self.draw_page_decorations(num_pages)
                super().showPage()
            super().save()

        def draw_page_decorations(self, page_count: int):
            self.saveState()
            self.setFont("Helvetica", 8)
            self.setFillColor(colors.HexColor("#64748b"))

            # Running header on page 2+
            if self._pageNumber > 1:
                self.drawString(40, letter[1] - 32, "CaseGraph Legal Intelligence — Official Case Report")
                self.setStrokeColor(colors.HexColor("#e2e8f0"))
                self.setLineWidth(0.5)
                self.line(40, letter[1] - 36, letter[0] - 40, letter[1] - 36)

            # Running footer on all pages
            page_str = f"Page {self._pageNumber} of {page_count}"
            self.drawRightString(letter[0] - 40, 25, page_str)
            self.drawString(40, 25, "CONFIDENTIAL — ATTORNEY WORK PRODUCT / PREPARED FOR COUNSEL")
            self.setStrokeColor(colors.HexColor("#cbd5e1"))
            self.setLineWidth(0.5)
            self.line(40, 35, letter[0] - 40, 35)

            self.restoreState()


class ReportAgent(BaseAgent[ReportInput, ReportOutput]):
    """
    Generates structured, professional case reports in multiple formats.
    """

    def __init__(self):
        super().__init__(
            name="ReportAgent",
            prompt_version="2.0.0",
            input_model=ReportInput,
            output_model=ReportOutput,
        )

    def _get_title(self, inp: ReportInput) -> str:
        type_titles = {
            "executive_summary": "Executive Summary Report",
            "review_summary": "E-Discovery Review Report",
            "privilege_log": "Privilege Log",
            "risk_assessment": "Comprehensive Risk Assessment Report",
        }
        return type_titles.get(inp.report_type, "E-Discovery Case Report")

    def _get_case_name(self, inp: ReportInput) -> str:
        if inp.case_info and inp.case_info.get("case_name"):
            return inp.case_info["case_name"]
        if inp.review_data and inp.review_data.get("case_name"):
            return inp.review_data["case_name"]
        return f"Case {inp.document_id}"

    def _to_json(self, inp: ReportInput) -> str:
        report = {
            "title": self._get_title(inp),
            "report_type": inp.report_type,
            "case_id": inp.document_id,
            "case_name": self._get_case_name(inp),
            "case_info": inp.case_info,
            "corpus_stats": inp.corpus_stats or inp.review_data,
            "key_metrics": {
                "documents_analyzed": (inp.corpus_stats or inp.review_data).get("total_documents", 0),
                "privileged_documents": (inp.corpus_stats or inp.review_data).get("privileged_documents", 0),
                "high_risk_flags": (inp.corpus_stats or inp.review_data).get("high_risk_documents")
                or (inp.corpus_stats or inp.review_data).get("high_risk_flags", 0),
            },
            "summary": {
                "total_documents": (inp.corpus_stats or inp.review_data).get("total_documents", 0),
            },
            "executive_summary": inp.executive_summary or "Case-level executive summary not yet available.",
            "executive_sections": inp.executive_sections,
            "summary_metadata": inp.summary_metadata,
            "top_entities": inp.top_entities,
            "findings": inp.findings,
            "timeline": inp.timeline,
            "privilege_log": inp.privilege_log,
            "risks": inp.risks,
            "recommendations": inp.recommendations,
            "generated_at": datetime.now(timezone.utc).isoformat(),
        }
        return json.dumps(report, indent=2, default=str)

    def _to_csv(self, inp: ReportInput) -> str:
        output = io.StringIO()
        writer = csv.writer(output)

        case_name = self._get_case_name(inp)
        writer.writerow(["Matter", case_name])
        writer.writerow(["Report Type", self._get_title(inp)])
        writer.writerow(["Generated At", datetime.now(timezone.utc).isoformat()])
        writer.writerow([])

        if inp.report_type == "privilege_log":
            writer.writerow([
                "Document Name",
                "Document ID",
                "Date",
                "Author / Custodian",
                "Privilege Type",
                "Legal Basis / Reason",
                "Review Status",
                "Withholding Status",
            ])
            for row in inp.privilege_log:
                writer.writerow([
                    row.get("document", ""),
                    row.get("document_id", ""),
                    row.get("date", ""),
                    row.get("author_custodian", ""),
                    row.get("privilege_type", ""),
                    row.get("legal_basis", ""),
                    row.get("review_status", ""),
                    row.get("withholding_status", ""),
                ])

        elif inp.report_type == "risk_assessment":
            writer.writerow(["Severity", "Category", "Explanation", "Source Document", "Document ID"])
            for r in inp.risks:
                writer.writerow([
                    r.get("severity") or r.get("risk_level", ""),
                    r.get("category", ""),
                    r.get("explanation") or r.get("reason", ""),
                    r.get("source_document", ""),
                    r.get("document_id", ""),
                ])

        elif inp.report_type == "executive_summary":
            writer.writerow(["Section", "Item", "Details"])
            writer.writerow(["Executive Summary", "Persisted Summary", inp.executive_summary or "Not available"])
            for ent in inp.top_entities:
                writer.writerow(["Key Entity", ent.get("name", ""), f"Type: {ent.get('type', '')}, Centrality: {ent.get('importance_score', '')}"])
            for ev in inp.timeline:
                writer.writerow(["Key Date", ev.get("date", ""), ev.get("description", "")])
            for f in inp.findings:
                writer.writerow(["Finding", f.get("title", ""), f.get("description", "")])

        else:  # review_summary
            writer.writerow(["Section", "Key", "Value"])
            stats = inp.corpus_stats or inp.review_data
            for k, v in stats.items():
                writer.writerow(["Corpus Statistics", k, v])
            for ev in inp.timeline:
                writer.writerow(["Timeline", ev.get("date", ""), ev.get("description", "")])
            for r in inp.risks:
                writer.writerow(["Risk", r.get("category", ""), f"{r.get('risk_level', '')}: {r.get('reason', '')}"])
            for rec in inp.recommendations:
                writer.writerow(["Recommendation", "", rec])

        return output.getvalue()

    def _to_markdown(self, inp: ReportInput) -> str:
        case_name = self._get_case_name(inp)
        title = self._get_title(inp)
        lines = [
            f"# {title}",
            f"**Matter:** {case_name}  ",
            f"**Case ID:** `{inp.document_id}`  ",
            f"**Generated:** {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}  ",
            "",
            "---",
            "",
        ]

        # Case info table
        if inp.case_info:
            lines.extend([
                "## Matter Overview",
                "",
                f"- **Matter Number:** {inp.case_info.get('matter_number', 'N/A')}",
                f"- **Client:** {inp.case_info.get('client', 'N/A')}",
                f"- **Case Type:** {inp.case_info.get('case_type', 'N/A')}",
                f"- **Status:** {inp.case_info.get('status', 'N/A')}",
                "",
            ])

        # Corpus Stats
        stats = inp.corpus_stats or inp.review_data
        if stats:
            lines.extend([
                "## Corpus Review Statistics",
                "",
                f"- **Total Documents:** {stats.get('total_documents', 0)}",
                f"- **Reviewed Documents:** {stats.get('reviewed_documents', stats.get('total_documents', 0))}",
                f"- **Flagged Documents:** {stats.get('flagged_documents', 0)}",
                f"- **Privileged Documents:** {stats.get('privileged_documents', 0)}",
                f"- **High Risk Documents:** {stats.get('high_risk_documents', stats.get('high_risk_flags', 0))}",
                "",
            ])

        if inp.report_type == "executive_summary":
            lines.extend([
                "## Persisted Executive Summary",
                "",
                inp.executive_summary or "Case-level executive summary not yet available.",
                "",
            ])
            if inp.executive_sections.get("legal"):
                lines.extend([
                    "### Legal Posture & Strategy",
                    "",
                    inp.executive_sections["legal"],
                    "",
                ])
            if inp.top_entities:
                lines.extend([
                    "## Key Entities & Roles",
                    "",
                    "| Entity Name | Type | Centrality Score |",
                    "| :--- | :--- | :--- |",
                ])
                for e in inp.top_entities[:10]:
                    score = f"{e.get('importance_score', 0):.4f}" if isinstance(e.get("importance_score"), (int, float)) else str(e.get("importance_score", "N/A"))
                    lines.append(f"| **{e.get('name', '')}** | {e.get('type', 'N/A')} | {score} |")
                lines.append("")

            if inp.findings:
                lines.extend([
                    "## Major Findings",
                    "",
                ])
                for f in inp.findings[:10]:
                    lines.append(f"- **[{f.get('severity', 'Info')}] {f.get('title', 'Finding')}:** {f.get('description', '')} _({f.get('source_document', '')})_")
                lines.append("")

        elif inp.report_type == "privilege_log":
            lines.extend([
                "## Formal Privilege Log",
                "",
            ])
            if inp.privilege_log:
                lines.extend([
                    "| Document | Doc ID | Date | Author / Custodian | Privilege Type | Legal Basis / Reason | Withholding |",
                    "| :--- | :--- | :--- | :--- | :--- | :--- | :--- |",
                ])
                for r in inp.privilege_log:
                    lines.append(
                        f"| {r.get('document', '')} | `{r.get('document_id', '')}` | {r.get('date', '')} | {r.get('author_custodian', '')} | {r.get('privilege_type', '')} | {r.get('legal_basis', '')} | {r.get('withholding_status', 'Withheld in Full')} |"
                    )
            else:
                lines.append("_No documents designated as privileged in this matter._")
            lines.append("")

        if inp.risks:
            lines.extend([
                "## Risk Assessment",
                "",
                "| Severity | Category | Explanation / Detail | Source Document |",
                "| :--- | :--- | :--- | :--- |",
            ])
            for r in inp.risks:
                sev = r.get("severity") or r.get("risk_level", "Medium")
                lines.append(
                    f"| **{sev}** | {r.get('category', 'Risk')} | {r.get('explanation') or r.get('reason', '')} | {r.get('source_document', '')} |"
                )
            lines.append("")
        elif inp.report_type == "risk_assessment":
            lines.extend([
                "## Risk Assessment Matrix",
                "",
                "_No elevated risks identified._",
                "",
            ])

        # Timeline (for executive and review summary)
        if inp.report_type in ("executive_summary", "review_summary"):
            lines.extend([
                "## Chronological Timeline",
                "",
            ])
            if inp.timeline:
                for ev in inp.timeline[:25]:
                    lines.append(f"- **{ev.get('date', 'N/A')}**: {ev.get('description', '')} _({ev.get('source_document', '')})_")
            else:
                lines.append("_No timeline events extracted._")
            lines.append("")

        if inp.recommendations:
            lines.extend(["## Recommendations", ""])
            for rec in inp.recommendations:
                lines.append(f"- {rec}")
            lines.append("")

        # Ensure compatibility with basic keyword searches in older tests
        if inp.report_type == "review_summary" and not inp.case_info:
            lines.insert(0, "# E-Discovery Review Report")

        return "\n".join(lines)

    def _to_pdf(self, inp: ReportInput) -> str:
        """Generates a professional legal PDF using ReportLab with NumberedCanvas."""
        if not REPORTLAB_AVAILABLE:
            raise AgentException("ReportLab is not installed or available.")

        buffer = io.BytesIO()
        doc = SimpleDocTemplate(
            buffer,
            pagesize=letter,
            leftMargin=40,
            rightMargin=40,
            topMargin=50,
            bottomMargin=50,
        )

        styles = getSampleStyleSheet()
        normal = styles["Normal"]

        # Custom typography
        title_style = ParagraphStyle(
            "DocTitle",
            parent=styles["Heading1"],
            fontName="Helvetica-Bold",
            fontSize=18,
            leading=22,
            textColor=colors.HexColor("#0f172a"),
            spaceAfter=4,
        )
        subtitle_style = ParagraphStyle(
            "DocSubtitle",
            parent=normal,
            fontName="Helvetica-Bold",
            fontSize=11,
            leading=15,
            textColor=colors.HexColor("#475569"),
            spaceAfter=12,
        )
        h2_style = ParagraphStyle(
            "SectionH2",
            parent=styles["Heading2"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=16,
            textColor=colors.HexColor("#1e3a8a"),
            spaceBefore=14,
            spaceAfter=6,
        )
        body_style = ParagraphStyle(
            "BodyP",
            parent=normal,
            fontName="Helvetica",
            fontSize=9,
            leading=13,
            textColor=colors.HexColor("#334155"),
            spaceAfter=6,
        )
        meta_label_style = ParagraphStyle(
            "MetaLabel",
            parent=normal,
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#475569"),
        )
        meta_val_style = ParagraphStyle(
            "MetaVal",
            parent=normal,
            fontName="Helvetica",
            fontSize=8,
            leading=11,
            textColor=colors.HexColor("#0f172a"),
        )
        th_style = ParagraphStyle(
            "TableHeader",
            parent=normal,
            fontName="Helvetica-Bold",
            fontSize=8,
            leading=10,
            textColor=colors.HexColor("#0f172a"),
        )
        tc_style = ParagraphStyle(
            "TableCell",
            parent=normal,
            fontName="Helvetica",
            fontSize=7.5,
            leading=9.5,
            textColor=colors.HexColor("#1e293b"),
        )

        def _esc(val: Any) -> str:
            if val is None:
                return ""
            return html.escape(str(val))

        story = []
        case_name = self._get_case_name(inp)
        report_title = self._get_title(inp)

        # Header Block
        story.append(Paragraph(_esc(report_title).upper(), title_style))
        story.append(Paragraph(f"MATTER: {_esc(case_name)}", subtitle_style))
        story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor("#1e3a8a"), spaceAfter=10))

        # Metadata & Case Overview Table
        info = inp.case_info or {}
        stats = inp.corpus_stats or inp.review_data or {}
        meta_data = [
            [
                Paragraph("<b>Matter Number:</b>", meta_label_style),
                Paragraph(_esc(info.get("matter_number") or info.get("case_number") or "N/A"), meta_val_style),
                Paragraph("<b>Total Documents:</b>", meta_label_style),
                Paragraph(_esc(stats.get("total_documents", 0)), meta_val_style),
            ],
            [
                Paragraph("<b>Client:</b>", meta_label_style),
                Paragraph(_esc(info.get("client", "N/A")), meta_val_style),
                Paragraph("<b>Reviewed Documents:</b>", meta_label_style),
                Paragraph(_esc(stats.get("reviewed_documents", stats.get("total_documents", 0))), meta_val_style),
            ],
            [
                Paragraph("<b>Case Type:</b>", meta_label_style),
                Paragraph(_esc(info.get("case_type", "Litigation")), meta_val_style),
                Paragraph("<b>Privileged Documents:</b>", meta_label_style),
                Paragraph(_esc(stats.get("privileged_documents", 0)), meta_val_style),
            ],
            [
                Paragraph("<b>Status:</b>", meta_label_style),
                Paragraph(_esc(info.get("status", "Active")).upper(), meta_val_style),
                Paragraph("<b>High-Risk Flags:</b>", meta_label_style),
                Paragraph(_esc(stats.get("high_risk_documents", stats.get("high_risk_flags", 0))), meta_val_style),
            ],
        ]
        meta_table = Table(meta_data, colWidths=[110, 156, 120, 146])
        meta_table.setStyle(TableStyle([
            ("BACKGROUND", (0, 0), (-1, -1), colors.HexColor("#f8fafc")),
            ("BOX", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
            ("INNERGRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#e2e8f0")),
            ("TOPPADDING", (0, 0), (-1, -1), 3),
            ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ("LEFTPADDING", (0, 0), (-1, -1), 6),
            ("RIGHTPADDING", (0, 0), (-1, -1), 6),
        ]))
        story.append(meta_table)
        story.append(Spacer(1, 10))

        # Content based on report type
        if inp.report_type == "executive_summary":
            story.append(Paragraph("1. Executive Summary & Case Synthesis", h2_style))
            exec_text = inp.executive_summary or "Case-level executive summary not yet available."
            story.append(Paragraph(_esc(exec_text).replace("\n", "<br/>"), body_style))

            if inp.executive_sections.get("legal"):
                story.append(Paragraph("2. Legal Posture & Exposure Analysis", h2_style))
                story.append(Paragraph(_esc(inp.executive_sections["legal"]).replace("\n", "<br/>"), body_style))

            if inp.top_entities:
                story.append(Paragraph("3. Key Case Entities", h2_style))
                ent_rows = [[
                    Paragraph("<b>Entity Name</b>", th_style),
                    Paragraph("<b>Type</b>", th_style),
                    Paragraph("<b>Centrality</b>", th_style),
                ]]
                for ent in inp.top_entities[:8]:
                    c_val = f"{ent.get('importance_score', 0):.4f}" if isinstance(ent.get("importance_score"), (int, float)) else str(ent.get("importance_score", "N/A"))
                    ent_rows.append([
                        Paragraph(_esc(ent.get("name", "")), tc_style),
                        Paragraph(_esc(ent.get("type", "N/A")), tc_style),
                        Paragraph(_esc(c_val), tc_style),
                    ])
                ent_table = Table(ent_rows, colWidths=[270, 140, 122], repeatRows=1)
                ent_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]))
                story.append(ent_table)
                story.append(Spacer(1, 10))

            if inp.findings:
                story.append(Paragraph("4. Major Case Findings", h2_style))
                f_rows = [[
                    Paragraph("<b>Category / Title</b>", th_style),
                    Paragraph("<b>Severity</b>", th_style),
                    Paragraph("<b>Description / Provision</b>", th_style),
                    Paragraph("<b>Source</b>", th_style),
                ]]
                for f in inp.findings[:8]:
                    f_rows.append([
                        Paragraph(_esc(f.get("title", "Finding")), tc_style),
                        Paragraph(_esc(f.get("severity", "Info")), tc_style),
                        Paragraph(_esc(f.get("description", "")), tc_style),
                        Paragraph(_esc(f.get("source_document", "")), tc_style),
                    ])
                f_table = Table(f_rows, colWidths=[110, 60, 242, 120], repeatRows=1)
                f_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]))
                story.append(f_table)
                story.append(Spacer(1, 10))

        elif inp.report_type == "privilege_log":
            story.append(Paragraph("Formal Privilege Log", h2_style))
            story.append(Paragraph(
                "Under applicable rules of civil procedure and professional conduct, the following documents have been designated as privileged attorney-client communications or attorney work product and are withheld from general production.",
                body_style,
            ))
            story.append(Spacer(1, 6))

            if inp.privilege_log:
                pl_rows = [[
                    Paragraph("<b>Document</b>", th_style),
                    Paragraph("<b>ID</b>", th_style),
                    Paragraph("<b>Date</b>", th_style),
                    Paragraph("<b>Custodian</b>", th_style),
                    Paragraph("<b>Privilege Type</b>", th_style),
                    Paragraph("<b>Legal Basis</b>", th_style),
                    Paragraph("<b>Status</b>", th_style),
                ]]
                for row in inp.privilege_log:
                    pl_rows.append([
                        Paragraph(_esc(row.get("document", "")), tc_style),
                        Paragraph(_esc(row.get("document_id", "")), tc_style),
                        Paragraph(_esc(row.get("date", "")), tc_style),
                        Paragraph(_esc(row.get("author_custodian", "")), tc_style),
                        Paragraph(_esc(row.get("privilege_type", "")), tc_style),
                        Paragraph(_esc(row.get("legal_basis", "")), tc_style),
                        Paragraph(_esc(row.get("withholding_status", "Withheld")), tc_style),
                    ])
                pl_table = Table(pl_rows, colWidths=[100, 52, 55, 70, 75, 125, 55], repeatRows=1)
                pl_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                    ("LEFTPADDING", (0, 0), (-1, -1), 3),
                    ("RIGHTPADDING", (0, 0), (-1, -1), 3),
                ]))
                story.append(pl_table)
            else:
                story.append(Paragraph("<i>No documents designated as privileged in this matter.</i>", body_style))

        elif inp.report_type == "risk_assessment":
            story.append(Paragraph("Risk Assessment Matrix", h2_style))
            story.append(Paragraph(
                "Risk items identified through multi-tier intelligence analysis across contractual obligations, regulatory compliance, and potential civil liabilities.",
                body_style,
            ))
            story.append(Spacer(1, 6))

            if inp.risks:
                r_rows = [[
                    Paragraph("<b>Severity</b>", th_style),
                    Paragraph("<b>Category</b>", th_style),
                    Paragraph("<b>Reason / Explanation</b>", th_style),
                    Paragraph("<b>Source Document</b>", th_style),
                ]]
                for r in inp.risks:
                    r_rows.append([
                        Paragraph(f"<b>{_esc(r.get('severity') or r.get('risk_level', 'Medium'))}</b>", tc_style),
                        Paragraph(_esc(r.get("category", "General")), tc_style),
                        Paragraph(_esc(r.get("explanation") or r.get("reason", "")), tc_style),
                        Paragraph(_esc(r.get("source_document", "")), tc_style),
                    ])
                r_table = Table(r_rows, colWidths=[60, 100, 232, 140], repeatRows=1)
                r_table.setStyle(TableStyle([
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                    ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                    ("TOPPADDING", (0, 0), (-1, -1), 3),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
                ]))
                story.append(r_table)
            else:
                story.append(Paragraph("<i>No elevated risks flagged.</i>", body_style))

        else:  # review_summary
            story.append(Paragraph("1. Review Progress & Corpus Breakdown", h2_style))
            story.append(Paragraph(
                f"E-Discovery document review progress summary for {_esc(case_name)}. All metrics are derived deterministically from processed database records.",
                body_style,
            ))

            if inp.risks:
                story.append(Paragraph("2. Elevated Risk Observations", h2_style))
                for r in inp.risks[:6]:
                    cat = _esc(r.get("category", "Risk"))
                    lvl = _esc(r.get("severity") or r.get("risk_level", "Medium"))
                    reason = _esc(r.get("explanation") or r.get("reason", ""))
                    story.append(Paragraph(f"• <b>[{lvl}] {cat}:</b> {reason}", body_style))

        # Timeline section for executive summary and review summary
        if inp.report_type in ("executive_summary", "review_summary") and inp.timeline:
            story.append(Spacer(1, 8))
            story.append(Paragraph("Chronological Timeline", h2_style))
            t_rows = [[
                Paragraph("<b>Date</b>", th_style),
                Paragraph("<b>Description</b>", th_style),
                Paragraph("<b>Source</b>", th_style),
            ]]
            for ev in inp.timeline[:20]:
                t_rows.append([
                    Paragraph(_esc(ev.get("date", "N/A")), tc_style),
                    Paragraph(_esc(ev.get("description", "")), tc_style),
                    Paragraph(_esc(ev.get("source_document", "")), tc_style),
                ])
            t_table = Table(t_rows, colWidths=[75, 337, 120], repeatRows=1)
            t_table.setStyle(TableStyle([
                ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#f1f5f9")),
                ("GRID", (0, 0), (-1, -1), 0.5, colors.HexColor("#cbd5e1")),
                ("TOPPADDING", (0, 0), (-1, -1), 3),
                ("BOTTOMPADDING", (0, 0), (-1, -1), 3),
            ]))
            story.append(t_table)

        doc.build(story, canvasmaker=NumberedCanvas)
        pdf_bytes = buffer.getvalue()
        buffer.close()

        return base64.b64encode(pdf_bytes).decode("utf-8")

    def _to_docx(self, inp: ReportInput) -> str:
        """Generates a professional Word DOCX report."""
        if not DOCX_AVAILABLE:
            raise AgentException("python-docx is not installed or available.")

        doc = docx.Document()
        case_name = self._get_case_name(inp)
        report_title = self._get_title(inp)

        # Title
        p_title = doc.add_heading(report_title, level=0)
        doc.add_paragraph(f"Matter: {case_name}")
        doc.add_paragraph(f"Generated: {datetime.now(timezone.utc).strftime('%Y-%m-%d %H:%M:%S UTC')}")

        # Metadata table
        doc.add_heading("Matter Overview", level=1)
        info = inp.case_info or {}
        stats = inp.corpus_stats or inp.review_data or {}

        t_meta = doc.add_table(rows=4, cols=2)
        t_meta.style = "Table Grid"
        t_meta.rows[0].cells[0].text = f"Matter Number: {info.get('matter_number', 'N/A')}"
        t_meta.rows[0].cells[1].text = f"Total Documents: {stats.get('total_documents', 0)}"
        t_meta.rows[1].cells[0].text = f"Client: {info.get('client', 'N/A')}"
        t_meta.rows[1].cells[1].text = f"Reviewed Documents: {stats.get('reviewed_documents', stats.get('total_documents', 0))}"
        t_meta.rows[2].cells[0].text = f"Case Type: {info.get('case_type', 'Litigation')}"
        t_meta.rows[2].cells[1].text = f"Privileged Documents: {stats.get('privileged_documents', 0)}"
        t_meta.rows[3].cells[0].text = f"Status: {info.get('status', 'Active')}"
        t_meta.rows[3].cells[1].text = f"High Risk Documents: {stats.get('high_risk_documents', stats.get('high_risk_flags', 0))}"

        if inp.report_type == "executive_summary":
            doc.add_heading("Executive Summary", level=1)
            doc.add_paragraph(inp.executive_summary or "Case-level executive summary not yet available.")
            if inp.executive_sections.get("legal"):
                doc.add_heading("Legal Posture & Strategy", level=2)
                doc.add_paragraph(inp.executive_sections["legal"])

            if inp.top_entities:
                doc.add_heading("Key Entities", level=2)
                t_ent = doc.add_table(rows=1, cols=3)
                t_ent.style = "Table Grid"
                t_ent.rows[0].cells[0].text = "Entity"
                t_ent.rows[0].cells[1].text = "Type"
                t_ent.rows[0].cells[2].text = "Centrality"
                for ent in inp.top_entities[:10]:
                    row = t_ent.add_row().cells
                    row[0].text = str(ent.get("name", ""))
                    row[1].text = str(ent.get("type", ""))
                    row[2].text = f"{ent.get('importance_score', 0):.4f}" if isinstance(ent.get("importance_score"), (int, float)) else str(ent.get("importance_score", "N/A"))

        elif inp.report_type == "privilege_log":
            doc.add_heading("Formal Privilege Log", level=1)
            if inp.privilege_log:
                t_pl = doc.add_table(rows=1, cols=7)
                t_pl.style = "Table Grid"
                hdr = t_pl.rows[0].cells
                hdr[0].text = "Document"
                hdr[1].text = "ID"
                hdr[2].text = "Date"
                hdr[3].text = "Custodian"
                hdr[4].text = "Privilege Type"
                hdr[5].text = "Legal Basis"
                hdr[6].text = "Withholding"
                for r in inp.privilege_log:
                    row = t_pl.add_row().cells
                    row[0].text = str(r.get("document", ""))
                    row[1].text = str(r.get("document_id", ""))
                    row[2].text = str(r.get("date", ""))
                    row[3].text = str(r.get("author_custodian", ""))
                    row[4].text = str(r.get("privilege_type", ""))
                    row[5].text = str(r.get("legal_basis", ""))
                    row[6].text = str(r.get("withholding_status", "Withheld"))
            else:
                doc.add_paragraph("No documents designated as privileged in this matter.")

        elif inp.report_type == "risk_assessment":
            doc.add_heading("Risk Assessment Matrix", level=1)
            if inp.risks:
                t_r = doc.add_table(rows=1, cols=4)
                t_r.style = "Table Grid"
                hdr = t_r.rows[0].cells
                hdr[0].text = "Severity"
                hdr[1].text = "Category"
                hdr[2].text = "Explanation"
                hdr[3].text = "Source Document"
                for r in inp.risks:
                    row = t_r.add_row().cells
                    row[0].text = str(r.get("severity") or r.get("risk_level", "Medium"))
                    row[1].text = str(r.get("category", "General"))
                    row[2].text = str(r.get("explanation") or r.get("reason", ""))
                    row[3].text = str(r.get("source_document", ""))
            else:
                doc.add_paragraph("No elevated risks flagged.")

        if inp.timeline and inp.report_type in ("executive_summary", "review_summary"):
            doc.add_heading("Chronological Timeline", level=1)
            t_time = doc.add_table(rows=1, cols=3)
            t_time.style = "Table Grid"
            t_time.rows[0].cells[0].text = "Date"
            t_time.rows[0].cells[1].text = "Description"
            t_time.rows[0].cells[2].text = "Source Document"
            for ev in inp.timeline[:25]:
                row = t_time.add_row().cells
                row[0].text = str(ev.get("date", "N/A"))
                row[1].text = str(ev.get("description", ""))
                row[2].text = str(ev.get("source_document", ""))

        buffer = io.BytesIO()
        doc.save(buffer)
        docx_bytes = buffer.getvalue()
        buffer.close()

        return base64.b64encode(docx_bytes).decode("utf-8")

    async def _execute(self, validated_input: ReportInput) -> dict:
        logger.info(f"[{self.name}] Generating {validated_input.format} {validated_input.report_type} report for {validated_input.document_id}")

        fmt = validated_input.format.lower()
        generators = {
            "markdown": self._to_markdown,
            "json": self._to_json,
            "csv": self._to_csv,
            "pdf": self._to_pdf,
            "docx": self._to_docx,
        }

        generator = generators.get(fmt)
        if not generator:
            raise AgentException(f"Unsupported report format: {fmt}")

        content = generator(validated_input)

        return {
            "document_id": validated_input.document_id,
            "report_type": validated_input.report_type,
            "format": fmt,
            "content": content,
            "confidence_score": 1.0,
        }
